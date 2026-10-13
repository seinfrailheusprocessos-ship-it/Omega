"""Gera levantamento_parametrico.xlsx (com formulas) a partir de etapa_b/resultado.json (MP) e etapa_b/resultado_justica.json (Justiça)."""
import json, os, sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

S, OUT = sys.argv[1], sys.argv[2]
R = json.load(open(os.path.join(S, 'etapa_b', 'resultado.json')))
RJ = json.load(open(os.path.join(S, 'etapa_b', 'resultado_justica.json')))
assert RJ['BDI'] == R['BDI'] and RJ['CONT'] == R['CONT'] and RJ['FREG'] == R['FREG']
for o in R['objetos']: o['levantamento'] = 'MP'
OBJS = R['objetos'] + RJ['objetos']
PROV = 'APROVADO em 13/10/2026'
F = 'Arial'
fn = lambda **k: Font(name=F, size=k.pop('size', 10), **k)
BLUE, GREEN, RED = '0000FF', '008000', 'C00000'
YEL = PatternFill('solid', fgColor='FFFF00')
HEAD = PatternFill('solid', fgColor='D9E1F2')
TOT = PatternFill('solid', fgColor='F2F2F2')
thin = Side(style='thin', color='BFBFBF'); BOX = Border(left=thin, right=thin, top=thin, bottom=thin)
MON, QTD, PCT = '#,##0.00;(#,##0.00);"-"', '#,##0.00;(#,##0.00);"-"', '0.00%'
WRAP = Alignment(wrap_text=True, vertical='top')

def short(n, k):
    if k.startswith('JUSTICA_TRABALHO'): s = 'TRT_' + k.split('_')[-1][:16]
    elif k.startswith('JUSTICA_'): s = ('TRF1_' if '.4.01.' in k else 'TJBA_') + k[8:23]
    elif k.startswith('TCM_'): s = k.split('-')[0]
    elif k.startswith('SN_MPF'): s = 'SN_MPF_FNDE_2023'
    elif k.startswith('SN_'): s = 'SN_CAMPO_DO_RODAO'
    elif k.startswith('MPF_'): s = 'MPF_' + k[12:]
    elif k.startswith('MPT_'): s = 'MPT_' + k[4:15]
    else: s = k
    return f'{n:02d}_{s}'[:31]

def titulo(ws, txt, sub=None):
    ws['A1'] = txt; ws['A1'].font = fn(bold=True, size=12)
    ws['A2'] = 'Parâmetros e triagem ' + PROV; ws['A2'].font = fn(bold=True, color=GREEN)
    if sub: ws['A3'] = sub; ws['A3'].font = fn(italic=True)

wb = Workbook()
wsR = wb.active; wsR.title = 'RESUMO'
wsP = wb.create_sheet('PARAMETROS')

# ------------------------------------------------------------------ PARAMETROS
titulo(wsP, 'Parâmetros da estimativa (BDI, contingência e fator regional)',
       'Custo = Σ(Quantidade × Preço unitário) × (1 + BDI) × (1 + Contingência) × Fator regional. Células amarelas: valores aprovados em 13/10/2026 (podem ser alterados; a planilha recalcula).')
hdr = ['Parâmetro', 'Mínimo', 'Provável', 'Máximo', 'Fonte / justificativa']
for j, h in enumerate(hdr, 1):
    c = wsP.cell(5, j, h); c.font = fn(bold=True); c.fill = HEAD; c.border = BOX
par = [('BDI - edificações, encostas, quadras, cercamento e outros', R['BDI']['edif'], 'Acórdão n. 2.622/2013-TCU-Plenário, "construção de edifícios": valor médio, único nos três cenários (sem CPRB; regime não desonerado).'),
       ('BDI - pavimentação, drenagem e infraestrutura de vias', R['BDI']['via'], 'Acórdão n. 2.622/2013-TCU-Plenário, "construção de rodovias e ferrovias": valor médio, único nos três cenários (sem CPRB).'),
       ('Contingência (projeto, sondagem, topografia, licenciamento, fiscalização)', R['CONT'], 'Parâmetro técnico declarado; teto de referência: precisão de ±30% para estudos preliminares (OT-IBR 004/2012, IBRAOP).'),
       ('Fator regional', R['FREG'], '1,00 nos três cenários: SINAPI-BA, localidade Salvador, regime não desonerado, adotado como referência de preço para Ilhéus (decisão de 13/10/2026).')]
for i, (nome, v, fonte) in enumerate(par, 6):
    wsP.cell(i, 1, nome).font = fn()
    for j, c in enumerate(('min', 'prov', 'max'), 2):
        x = wsP.cell(i, j, v[c]); x.font = fn(color=BLUE); x.fill = YEL; x.number_format = PCT if i < 9 else '0.00'
        x.border = BOX
    wsP.cell(i, 5, fonte).font = fn(); wsP.cell(i, 5).alignment = WRAP
    wsP.cell(i, 1).border = BOX; wsP.cell(i, 5).border = BOX
wsP['A11'] = 'Preços unitários: SINAPI/CAIXA, ref. 08/2026, Bahia (localidade Salvador), regime não desonerado; encargos sociais horista 117,73% e mensalista 72,85%.'
wsP['A12'] = 'Ver proposta_bdi_contingencia_fator.md para as fontes conferidas e as limitações da conferência.'
for a in ('A11', 'A12'): wsP[a].font = fn(italic=True)
for col, w in zip('ABCDE', (62, 12, 12, 12, 95)): wsP.column_dimensions[col].width = w
PREF = {('edif', 'min'): 'PARAMETROS!$B$6', ('edif', 'prov'): 'PARAMETROS!$C$6', ('edif', 'max'): 'PARAMETROS!$D$6',
        ('via', 'min'): 'PARAMETROS!$B$7', ('via', 'prov'): 'PARAMETROS!$C$7', ('via', 'max'): 'PARAMETROS!$D$7',
        ('cont', 'min'): 'PARAMETROS!$B$8', ('cont', 'prov'): 'PARAMETROS!$C$8', ('cont', 'max'): 'PARAMETROS!$D$8',
        ('freg', 'min'): 'PARAMETROS!$B$9', ('freg', 'prov'): 'PARAMETROS!$C$9', ('freg', 'max'): 'PARAMETROS!$D$9'}

# ------------------------------------------------------------------ abas por objeto
links = []
for o in OBJS:
    nome = short(o['n'], o['grupo_objeto'])
    ws = wb.create_sheet(nome)
    titulo(ws, f"{o['n']:02d} - {o['titulo']}")
    info = [('Grupo-objeto', o['grupo_objeto']), ('Levantamento', o['levantamento']), ('Origem', o['origem']), ('Categoria', o['categoria']), ('Local', o['local'] or 'não informado'),
            ('Fonte das quantidades', o['fonte_quant']), ('Preços', 'SINAPI/CAIXA ref. 08/2026, BA, não desonerado (código de cada item na coluna B)'),
            ('BDI aplicado', 'edificações (Acórdão 2.622/2013, construção de edifícios)' if o['bdi_tipo'] == 'edif' else 'vias e drenagem (Acórdão 2.622/2013, rodovias e ferrovias)')]
    r = 4
    for a, b in info:
        ws.cell(r, 1, a).font = fn(bold=True); ws.cell(r, 2, b).font = fn(); r += 1
    ws.cell(r, 1, 'Premissas').font = fn(bold=True)
    for p in o['premissas']:
        ws.cell(r, 2, p).font = fn(); ws.cell(r, 2).alignment = Alignment(wrap_text=False); r += 1
    r += 1
    H = ['Macroetapa', 'Código SINAPI', 'Tipo', 'Descrição', 'Unid.', 'Preço unit. (R$)', 'Qtd. mínimo', 'Qtd. provável', 'Qtd. máximo',
         'Total mínimo (R$)', 'Total provável (R$)', 'Total máximo (R$)', 'Memória de cálculo']
    for j, h in enumerate(H, 1):
        c = ws.cell(r, j, h); c.font = fn(bold=True); c.fill = HEAD; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='center')
    r0 = r + 1; r = r0
    for it in o['itens']:
        vals = [it['macro'], it['codigo'], it['tipo'], it['descricao'], it['unidade'], it['preco'],
                round(it['q']['min'], 2), round(it['q']['prov'], 2), round(it['q']['max'], 2)]
        for j, v in enumerate(vals, 1):
            c = ws.cell(r, j, v); c.font = fn(color=BLUE) if j in (6, 7, 8, 9) else fn(); c.border = BOX
            if j in (6, 7, 8, 9): c.number_format = MON if j == 6 else QTD
        for j, q in zip((10, 11, 12), 'GHI'):
            c = ws.cell(r, j, f'=F{r}*{q}{r}'); c.number_format = MON; c.font = fn(); c.border = BOX
        ws.cell(r, 13, it['memoria']).font = fn(); ws.cell(r, 13).border = BOX
        r += 1
    r1 = r - 1
    rd, rb, rc, rf, rt = r, r + 1, r + 2, r + 3, r + 4
    ws.cell(rd, 9, 'Custo direto').font = fn(bold=True)
    ws.cell(rb, 9, 'BDI').font = fn(bold=True)
    ws.cell(rc, 9, 'Contingência').font = fn(bold=True)
    ws.cell(rf, 9, 'Fator regional').font = fn(bold=True)
    ws.cell(rt, 9, 'TOTAL ESTIMADO (R$)').font = fn(bold=True)
    for j, c in zip((10, 11, 12), ('min', 'prov', 'max')):
        L = get_column_letter(j)
        x = ws.cell(rd, j, f'=SUM({L}{r0}:{L}{r1})'); x.number_format = MON; x.font = fn(bold=True)
        x = ws.cell(rb, j, f"={PREF[(o['bdi_tipo'], c)]}"); x.number_format = PCT; x.font = fn(color=GREEN)
        x = ws.cell(rc, j, f"={PREF[('cont', c)]}"); x.number_format = PCT; x.font = fn(color=GREEN)
        x = ws.cell(rf, j, f"={PREF[('freg', c)]}"); x.number_format = '0.00'; x.font = fn(color=GREEN)
        x = ws.cell(rt, j, f'={L}{rd}*(1+{L}{rb})*(1+{L}{rc})*{L}{rf}'); x.number_format = MON; x.font = fn(bold=True); x.fill = TOT
    ws.cell(rt + 2, 1, 'Estimativa preliminar paramétrica, não vinculante, para planejamento orçamentário. Não substitui o orçamento detalhado do projeto. Parâmetros ' + PROV + '.').font = fn(italic=True, color=RED)
    for col, w in zip('ABCDEFGHIJKLM', (34, 11, 17, 70, 7, 13, 12, 12, 12, 16, 16, 16, 60)): ws.column_dimensions[col].width = w
    ws.freeze_panes = ws.cell(r0, 5)
    links.append((o, nome, rt))

# ------------------------------------------------------------------ RESUMO
titulo(wsR, 'Levantamento paramétrico - demandas de MP, Justiça e TCM (SEINFRA Ilhéus)',
       'Estimativa preliminar paramétrica, não vinculante, para planejamento orçamentário. Preços SINAPI-BA 08/2026, não desonerado. Data: 13/10/2026. Levantamentos MP (pacote de 12/10) e Justiça (pacote de 13/10), sem dupla contagem.')
H = ['Nº', 'Grupo-objeto', 'Objeto', 'Origem', 'Levantamento', 'Categoria', 'Fonte das quantidades', 'Mínimo (R$)', 'Provável (R$)', 'Máximo (R$)', 'Aba']
for j, h in enumerate(H, 1):
    c = wsR.cell(5, j, h); c.font = fn(bold=True); c.fill = HEAD; c.border = BOX; c.alignment = Alignment(wrap_text=True, vertical='center')
r = 6
for o, nome, rt in links:
    vals = [o['n'], o['grupo_objeto'], o['titulo'], o['origem'], o['levantamento'], o['categoria'], o['fonte_quant']]
    for j, v in enumerate(vals, 1):
        c = wsR.cell(r, j, v); c.font = fn(); c.border = BOX
    for j, L in zip((8, 9, 10), 'JKL'):
        c = wsR.cell(r, j, f"='{nome}'!{L}{rt}"); c.number_format = MON; c.font = fn(color=GREEN); c.border = BOX
    wsR.cell(r, 11, nome).font = fn(); wsR.cell(r, 11).border = BOX
    r += 1
ra, rz = 6, r - 1
wsR.cell(r, 7, 'TOTAL GERAL').font = fn(bold=True)
for j, L in zip((8, 9, 10), 'HIJ'):
    c = wsR.cell(r, j, f'=SUM({L}{ra}:{L}{rz})'); c.number_format = MON; c.font = fn(bold=True); c.fill = TOT
rg = r
def bloco(r, titulo_b, col_crit, chaves, rotulo_soma='Soma (confere com o total geral)'):
    wsR.cell(r, 1, titulo_b).font = fn(bold=True); r += 1
    for j, h in enumerate(['Chave', '', '', '', '', '', 'Objetos', 'Mínimo (R$)', 'Provável (R$)', 'Máximo (R$)'], 1):
        if h: c = wsR.cell(r, j, h); c.font = fn(bold=True); c.fill = HEAD
    r += 1; r_ini = r
    for ch in chaves:
        wsR.cell(r, 1, ch).font = fn()
        wsR.cell(r, 7, f'=COUNTIF(${col_crit}${ra}:${col_crit}${rz},A{r})').font = fn()
        for j, L in zip((8, 9, 10), 'HIJ'):
            c = wsR.cell(r, j, f'=SUMIF(${col_crit}${ra}:${col_crit}${rz},$A{r},{L}${ra}:{L}${rz})'); c.number_format = MON; c.font = fn()
        r += 1
    wsR.cell(r, 1, rotulo_soma).font = fn(italic=True, bold=rotulo_soma.startswith('GERAL'))
    for j, L in zip((8, 9, 10), 'HIJ'):
        c = wsR.cell(r, j, f'=SUM({L}{r_ini}:{L}{r-1})'); c.number_format = MON; c.font = fn(italic=True, bold=rotulo_soma.startswith('GERAL'))
        if rotulo_soma.startswith('GERAL'): c.fill = TOT
    return r + 2
r = bloco(rg + 2, 'TOTAIS MP, JUSTIÇA E GERAL (MP = pacote do MP de 12/10, inclui processos judiciais e TCM nele; Justiça = objetos novos do pacote da Justiça de 13/10)',
          'E', ['MP', 'Justiça'], 'GERAL (MP + Justiça, sem dupla contagem)')
# paramétricos da Justiça já calculados no MP: só referência (já somados no MP)
wsR.cell(r, 1, 'Objetos "orçamento paramétrico" da Justiça já calculados no MP (não somados de novo; valor da aba do MP)').font = fn(bold=True); r += 1
abas = {o['grupo_objeto']: (nome, rt) for o, nome, rt in links}
for x in RJ['ja_calculados_no_mp']:
    nome_mp, rt_mp = abas[x['ja_no_mp']]
    wsR.cell(r, 2, x['grupo_objeto']).font = fn(); wsR.cell(r, 3, next(o['titulo'] for o, _, _ in links if o['grupo_objeto'] == x['ja_no_mp'])).font = fn(); wsR.cell(r, 7, 'no MP: ' + x['ja_no_mp']).font = fn(italic=True)
    for j, L in zip((8, 9, 10), 'JKL'):
        c = wsR.cell(r, j, f"='{nome_mp}'!{L}{rt_mp}"); c.number_format = MON; c.font = fn(color=GREEN, italic=True)
    wsR.cell(r, 11, nome_mp).font = fn(italic=True); r += 1
r += 1
r = bloco(r, 'Total por origem', 'D', ['MPBA', 'MPF', 'MPT', 'Justiça', 'TCM', 'outro'])
cats = []
for o, _, _ in links:
    if o['categoria'] not in cats: cats.append(o['categoria'])
r = bloco(r, 'Total por categoria', 'F', cats)
notas = ['Leitura: o valor "provável" é a referência para reserva orçamentária; mínimo e máximo dão a faixa de incerteza.',
         'Quantidades: metragem do documento quando existe; nos demais casos, demanda-tipo com premissas na aba de cada objeto (a confirmar em campo/projeto).',
         'BDI, contingência e fator regional: aba PARAMETROS (' + PROV + ').',
         'Origem "outro" = processo interno "Resposta Ministério Público" sem indicação do ramo do MP (somado em MP).',
         'Justiça: só os objetos "orçamento paramétrico" do pacote da Justiça que não estavam calculados no MP (triagem_provisoria_justica.json, aprovada). Retirados por decisão do usuário: reforma do Renascer e Canal do Malhado (licitação própria) e projeto de incêndio da Casa Lar Masculina (tratados à parte).',
         'Não usar como orçamento de licitação. Estimativa não vinculante, com SINAPI como base oficial de referência, nos termos da Lei n. 14.133/2021 e do Decreto Municipal n. 269/2025.']
for t in notas:
    wsR.cell(r, 1, t).font = fn(italic=True); r += 1
for col, w in zip('ABCDEFGHIJK', (5, 36, 62, 9, 10, 20, 34, 16, 16, 16, 26)): wsR.column_dimensions[col].width = w
wsR.freeze_panes = 'C6'

# ------------------------------------------------------------------ DEMANDAS-TIPO e FONTES
wsD = wb.create_sheet('DEMANDAS-TIPO')
titulo(wsD, 'Demandas-tipo (premissas aprovadas em 13/10/2026; confirmar quantidades em campo ou projeto)')
DT = [('Via urbana com drenagem e pavimentação', 'metro de via', 'Pista 5/6/7 m; passeios 1,2/1,5/2,0 m; base+sub-base 15/20/30 cm; CBUQ 4/5/5 cm; rede DN 600/600/800; PV a cada 40 m; 2 BL a cada 40 m; meio-fio, sarjeta, passeio, rampas e faixas.'),
      ('Microdrenagem em via existente', 'metro de rede', 'Vala DN+0,6 x DN+1,0 m; tubo de concreto; PV a cada 40 m; 2 BL a cada 40 m com ligação DN 400; recomposição de base 20 cm e CBUQ 5 cm na faixa da vala; bota-fora a 10 km.'),
      ('Contenção de encosta (solo grampeado + concreto projetado)', 'm² de face', 'Grampos em malha 2,0/1,5/1,25 m, comprimento 6/8/12 m; concreto projetado 10/10/15 cm; drenos 0,6 m/m²; canaletas de crista e pé; descidas d\'água; grama em 30% da área. Calibrada com a média do Novo PAC Encostas citada nos documentos (R$ 1.692/m²).'),
      ('Edificação térrea nova', 'm² construído', 'Concreto 0,12 m³/m²; aço 10 kg/m²; fôrma 1,1 m²/m²; alvenaria 1,5 m²/m²; cobertura cerâmica 1,25 m²/m²; forro PVC; piso cerâmico; portas 1/15 m²; janelas 0,15 m²/m²; elétrica e hidrossanitária simplificadas; prevenção contra incêndio básica. Remanescente de obra = fração da área.'),
      ('Reforma predial', 'm² reformado', 'Fração refeita de cobertura/forro/piso/reboco/esquadrias/instalações: 20/20/10/10/10/20% (leve), 50/50/30/30/30/50% (média), 100/100/60/60/60/100% (pesada); pintura integral; extintores, emergência e rota de fuga.'),
      ('Quadra poliesportiva coberta', 'unidade', 'Piso 32 x 20 m em concreto armado 8 cm; cobertura metálica 748 m² (18/25/32 kg/m² de aço); alambrado 104 x 2 m; tabelas; refletores; vestiário 0/40/60 m².'),
      ('Campo de futebol', 'unidade', 'Campo 70x45 / 90x60 / 100x64 m; regularização; drenos a cada 5 m; grama em placas; alambrado; calçadas; vestiário 60/80/120 m².'),
      ('Abrigo de passageiros', 'unidade', 'Abrigo 4 x 3 m: fundação, estrutura de aço 300 kg, telha termoacústica 12 m², piso 12 m², banco e piso podotátil.'),
      ('Adequação de cozinha escolar', 'unidade', 'Abrigo de GLP, janela de ventilação, 15 m² de revestimento de parede e de piso, caixa de gordura, bancada, pontos elétricos, luminárias e extintor (sem exaustor: sem preço no SINAPI-BA).'),
      ('Banheiro com chuveiro + grade', 'unidade', 'Banheiro de cerca de 4 m² com louças, chuveiro, instalações e porta; grade de aço de 2,1 m² na porta de entrada.'),
      ('Proteção emergencial de fachada', 'verba por quantidades', 'Tapume, estroncas de eucalipto com mão de obra, montagem de andaime fachadeiro, tela de proteção e lona na cobertura.'),
      ('Adequação de segurança de estádio (TAC)', 'verba por quantidades', 'Só itens de segurança do TAC: extintores PQS, água e CO2, luminárias e placas de emergência, barras antipânico nas saídas, guarda-corpos, SPDA (captores, cordoalha, hastes e caixas) e quadros com DR nas casas de força. Sem reforma, pintura ou acessibilidade; sem alarme (sem preço no SINAPI-BA).'),
      ('Casa térrea + parque linear', 'unidade / m²', 'Casa de 40/42/45 m² com a edificação-tipo (perfil residencial); parque com gramado em 70% e passeios em 20% da área.'),
      ('Infraestrutura básica de loteamento (Justiça)', 'metro de via', 'Via urbana-tipo + rede de água PVC DN 50 (vala 0,6 x 1,0 m) + rede de esgoto PVC DN 150 (vala 0,8 x 1,5 m, PV pré-moldado a cada 60 m) + luminária LED 68-97 W com braço a cada 35 m, em poste existente.'),
      ('Cercamento de área (Justiça)', 'metro de perímetro', 'Alambrado de 2 m com tela de arame galvanizado, mourões e mureta de concreto, concertina no topo, limpeza de faixa de 3 m, placas a cada 100 m e portões em gradil de aço de 4,0 x 2,2 m.')]
for j, h in enumerate(['Demanda-tipo', 'Unidade-chave', 'Premissas e coeficientes (mín/prov/máx)'], 1):
    c = wsD.cell(4, j, h); c.font = fn(bold=True); c.fill = HEAD
for i, row in enumerate(DT, 5):
    for j, v in enumerate(row, 1):
        c = wsD.cell(i, j, v); c.font = fn(); c.alignment = WRAP
for col, w in zip('ABC', (40, 18, 140)): wsD.column_dimensions[col].width = w

wsF = wb.create_sheet('FONTES')
titulo(wsF, 'Fontes e limitações')
FT = ['SINAPI/CAIXA - Relatórios mensais, ref. 08/2026 (publicado em 11/09/2026), recorte BA, regime sem desoneração: composições (aba CSD) e insumos (aba ISD). Arquivo: _PRECOS/SINAPI_BA_NAO_DESONERADO_2026_08.zip.',
      'Acórdão n. 2.622/2013-TCU-Plenário - faixas referenciais de BDI (construção de edifícios; construção de rodovias e ferrovias). Conferido em reproduções públicas do texto.',
      'IBRAOP - OT-IBR 004/2012 (precisão do orçamento de obras públicas): ±30% para estudos preliminares em edificações. Conferido em documentos que reproduzem a OT.',
      'Parâmetros (BDI, contingência e fator regional) e triagens aprovados pelo usuário em 13/10/2026.',
      'CAIXA - SINAPI: metodologias e conceitos (coleta de preços nas capitais, sem frete). Conferido em reprodução pública.',
      'Base legal: nos termos da Lei n. 14.133/2021 e do Decreto Municipal n. 269/2025 (citados sem artigo).',
      'Quantidades: documentos dos pacotes MP_FILTRADO de 12/10/2026 (mp_objetos.json) e JUSTICA_FILTRADO de 13/10/2026 (justica_objetos.json) e demandas-tipo desta planilha.',
      'Limitação: composições sem preço na BA (ex.: imprimação, pintura de ligação, exaustores, central de alarme, acionadores e sirenes, locação de andaime) não estão incluídas; quantidades da maioria dos objetos são premissas a confirmar.',
      'Limitação: OpenStreetMap e sites oficiais não puderam ser acessados neste ambiente; comprimentos de ruas não foram medidos.',
      'Contratos da SEINFRA e editais de outras prefeituras não foram usados como base de preço.']
for i, t in enumerate(FT, 4):
    c = wsF.cell(i, 1, t); c.font = fn(); c.alignment = WRAP
wsF.column_dimensions['A'].width = 160

for ws in wb.worksheets:
    ws.sheet_view.showGridLines = True
    ws.page_setup.orientation = 'landscape'
wb.save(OUT)
print('ok', OUT, len(wb.worksheets), 'abas')
