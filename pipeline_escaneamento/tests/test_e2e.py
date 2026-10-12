"""Teste de ponta a ponta com documentos sintéticos (OCR real). Mede separação, categoria, nome e soltos."""
import csv
import json
import os

import pymupdf
import pytest
from conftest import precisa_ocr

from pipeline.avaliar import avaliar, formatar
from pipeline.util import P, sha256_arquivo

pytestmark = precisa_ocr


def test_resumo_sem_erro_fatal(e2e):
    assert e2e["resumo"]["ok"], e2e["resumo"]["mensagem"]


def test_metricas_de_qualidade(e2e):
    m = avaliar(e2e["itens"], e2e["verdade"])
    print("\n" + formatar(m))
    n = m["docs_verdade"]
    assert m["sep_ok"] / n >= 0.80, "separação correta"
    assert m["cat_ok"] / n >= 0.75, "categoria correta"
    assert m["nome_ok"] / n >= 0.75, "nome final correto"
    assert m["soltos_acertados"] >= m["soltos_esperados"] - 1, "páginas soltas tratadas"
    assert m["brancas_acertadas"] == m["brancas_esperadas"], "páginas em branco registradas"


def test_originais_nunca_alterados(e2e):
    for nome, h in e2e["hashes_antes"].items():
        assert sha256_arquivo(os.path.join(e2e["entrada"], nome)) == h, nome


def test_arquivos_problema_vao_para_erros(e2e):
    erros = [i for i in e2e["itens"] if i["status"] == "ERRO"]
    motivos = " ".join(i["motivo"] for i in erros)
    assert "corrompido" in motivos and "senha" in motivos
    pasta = e2e["pastas"].erros
    nomes = os.listdir(P(pasta))
    assert "corrompido.pdf" in nomes and "protegido.pdf" in nomes
    assert any(n.endswith(".motivo.txt") for n in nomes)


def test_pdf_pesquisavel_e_nome_final_e_json_ao_lado(e2e):
    docs = [i for i in e2e["itens"] if i["status"] in ("AUTO", "REVISAR")]
    assert docs
    for it in docs[:6]:
        assert os.path.exists(P(it["caminho_final"])), it["caminho_final"]
        assert it["nome_final"] == it["nome_final"].upper().replace(".PDF", ".pdf")
        assert " " not in it["nome_final"] and "(" not in it["nome_final"]
        with pymupdf.open(str(P(it["caminho_final"]))) as d:
            assert d.page_count == it["n_paginas"]
            assert len(d[0].get_text().strip()) > 50, "PDF sem camada de texto"
        j = json.load(open(P(os.path.splitext(it["caminho_final"])[0] + ".json"), encoding="utf-8"))
        assert j["id"] == it["id"] and j["contrato_versao"] == 1 and "campos" in j
        assert os.path.exists(P(it["caminho_texto"]))


def test_originais_ocr_e_mapa(e2e):
    pasta = e2e["pastas"].originais_ocr
    nomes = os.listdir(P(pasta))
    pdfs = [n for n in nomes if n.endswith("_OCR.pdf")]
    mapas = [n for n in nomes if n.endswith(".mapa.json")]
    assert len(pdfs) == len(mapas) >= 3
    mapa = json.load(open(P(os.path.join(pasta, mapas[0])), encoding="utf-8"))
    assert mapa["paginas"] and all(p for p in mapa["paginas"])


def test_catalogo_csv_xlsx_jsonl(e2e):
    from openpyxl import load_workbook
    p = e2e["pastas"]
    raw = open(P(p.csv), "rb").read()
    assert raw.startswith(b"\xef\xbb\xbf"), "CSV precisa de BOM UTF-8"
    cab = raw.decode("utf-8-sig").splitlines()[0].split(";")
    for col in ["id", "arquivo_origem", "hash_origem", "paginas_origem", "n_paginas", "categoria", "tipo", "orgao", "numero",
                "data", "assunto", "destinatario", "signatario_nome", "signatario_cargo", "processo", "objeto",
                "confianca_separacao", "confianca_categoria", "motivo", "status", "caminho_final", "nome_final", "caminho_texto"]:
        assert col in cab, col
    wb = load_workbook(str(P(p.xlsx)))
    assert wb.sheetnames == ["Resumo", "Documentos", "Soltos", "Revisar", "Erros"]
    linhas = [json.loads(l) for l in open(P(p.jsonl), encoding="utf-8")]
    assert {l["status"] for l in linhas} >= {"AUTO", "SOLTO", "ERRO"}
    assert all("campos" in l for l in linhas)


def test_status_json(e2e):
    st = json.load(open(P(e2e["pastas"].status), encoding="utf-8"))
    assert st["em_andamento"] is False and st["percentual"] == 100.0
    assert sum(st["por_categoria"].values()) == e2e["resumo"]["documentos"]
    assert os.path.exists(P(e2e["pastas"].progresso))


def test_idempotencia_nao_repete(e2e):
    from pipeline.run import executar
    n_antes = len(os.listdir(P(e2e["pastas"].originais_ocr)))
    r = executar(e2e["cfg"], console=False)
    assert r["documentos"] == 0 and r["ignorados"] >= 4
    assert len(os.listdir(P(e2e["pastas"].originais_ocr))) == n_antes
