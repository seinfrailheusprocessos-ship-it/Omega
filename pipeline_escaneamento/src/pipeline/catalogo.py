"""Catálogo em SQLite (idempotência + contrato para máquinas) e exportações (xlsx, csv, jsonl)."""
from __future__ import annotations

import csv
import datetime as dt
import json
import os
import sqlite3

from . import VERSAO
from .util import P, escrever_atomico

COLUNAS = ["id", "arquivo_origem", "hash_origem", "paginas_origem", "n_paginas", "categoria", "tipo", "orgao",
           "numero", "data", "assunto", "destinatario", "signatario_nome", "signatario_cargo", "processo",
           "objeto", "confianca_separacao", "confianca_categoria", "motivo", "status", "caminho_final",
           "nome_final", "caminho_texto"]
EXTRAS = ["processo_tipo", "caminho_origem", "caminho_json", "run_id", "criado_em", "campos_json"]
TODAS = COLUNAS + EXTRAS

SCHEMA = f"""
CREATE TABLE IF NOT EXISTS runs(id INTEGER PRIMARY KEY AUTOINCREMENT, inicio TEXT, fim TEXT, modo TEXT,
    versao TEXT, manifesto TEXT, desfeito INTEGER DEFAULT 0, resumo TEXT);
CREATE TABLE IF NOT EXISTS arquivos(hash TEXT PRIMARY KEY, caminho TEXT, nome TEXT, tamanho INTEGER,
    n_paginas INTEGER, status TEXT, motivo TEXT, run_id INTEGER, quando TEXT);
CREATE TABLE IF NOT EXISTS itens(id INTEGER PRIMARY KEY AUTOINCREMENT,
    arquivo_origem TEXT, hash_origem TEXT, paginas_origem TEXT, n_paginas INTEGER, categoria TEXT, tipo TEXT,
    orgao TEXT, numero TEXT, data TEXT, assunto TEXT, destinatario TEXT, signatario_nome TEXT,
    signatario_cargo TEXT, processo TEXT, objeto TEXT, confianca_separacao REAL, confianca_categoria REAL,
    motivo TEXT, status TEXT, caminho_final TEXT, nome_final TEXT, caminho_texto TEXT,
    processo_tipo TEXT, caminho_origem TEXT, caminho_json TEXT, run_id INTEGER, criado_em TEXT, campos_json TEXT);
CREATE INDEX IF NOT EXISTS ix_itens_hash ON itens(hash_origem);
CREATE INDEX IF NOT EXISTS ix_itens_run ON itens(run_id);
"""


def agora() -> str:
    return dt.datetime.now().isoformat(timespec="seconds")


class Catalogo:
    def __init__(self, caminho_db: str, somente_leitura: bool = False):
        if caminho_db == ":memory:":
            self.con = sqlite3.connect(":memory:")
        else:
            d = P(caminho_db)
            d.parent.mkdir(parents=True, exist_ok=True)
            from .util import _caminho_simples
            self.con = sqlite3.connect(caminho_db if _caminho_simples(caminho_db) else str(d), timeout=60)
        self.con.row_factory = sqlite3.Row
        self.con.execute("PRAGMA journal_mode=WAL")
        self.con.executescript(SCHEMA)
        self.somente_leitura = somente_leitura

    # ---- runs
    def nova_run(self, modo: str, manifesto: str) -> int:
        c = self.con.execute("INSERT INTO runs(inicio,modo,versao,manifesto) VALUES(?,?,?,?)",
                             (agora(), modo, VERSAO, manifesto))
        self.con.commit()
        return c.lastrowid

    def fechar_run(self, run_id: int, resumo: dict):
        self.con.execute("UPDATE runs SET fim=?, resumo=? WHERE id=?",
                         (agora(), json.dumps(resumo, ensure_ascii=False), run_id))
        self.con.commit()

    def ultima_run_desfazivel(self):
        return self.con.execute(
            "SELECT * FROM runs WHERE desfeito=0 AND modo='processar' AND "
            "(EXISTS(SELECT 1 FROM itens WHERE run_id=runs.id) OR EXISTS(SELECT 1 FROM arquivos WHERE run_id=runs.id)) "
            "ORDER BY id DESC LIMIT 1").fetchone()

    # ---- arquivos
    def arquivo(self, hash_: str):
        return self.con.execute("SELECT * FROM arquivos WHERE hash=?", (hash_,)).fetchone()

    def registrar_arquivo(self, hash_, caminho, nome, tamanho, n_paginas, status, motivo, run_id):
        self.con.execute("INSERT OR REPLACE INTO arquivos VALUES(?,?,?,?,?,?,?,?,?)",
                         (hash_, caminho, nome, tamanho, n_paginas, status, motivo, run_id, agora()))

    # ---- itens
    def inserir_item(self, **kw) -> int:
        kw.setdefault("criado_em", agora())
        cols = [c for c in TODAS if c != "id" and c in kw]
        c = self.con.execute(f"INSERT INTO itens({','.join(cols)}) VALUES({','.join('?' * len(cols))})",
                             [kw[k] for k in cols])
        return c.lastrowid

    def atualizar_item(self, id_: int, **kw):
        cols = list(kw)
        self.con.execute(f"UPDATE itens SET {','.join(c + '=?' for c in cols)} WHERE id=?",
                         [kw[c] for c in cols] + [id_])

    def caminhos_usados(self) -> set[str]:
        return {r[0].lower() for r in self.con.execute("SELECT caminho_final FROM itens WHERE caminho_final<>''")}

    def itens(self, where: str = "", args=()) -> list[dict]:
        q = "SELECT * FROM itens " + (("WHERE " + where) if where else "") + " ORDER BY id"
        return [dict(r) for r in self.con.execute(q, args)]

    def apagar_run(self, run_id: int):
        self.con.execute("DELETE FROM itens WHERE run_id=?", (run_id,))
        self.con.execute("DELETE FROM arquivos WHERE run_id=?", (run_id,))
        self.con.execute("UPDATE runs SET desfeito=1 WHERE id=?", (run_id,))
        self.con.commit()

    def commit(self):
        self.con.commit()

    def rollback(self):
        self.con.rollback()

    def fechar(self):
        self.con.commit()
        self.con.close()


# ----------------------------------------------------------------- exportações
def _linha_publica(it: dict) -> dict:
    return {c: it.get(c) for c in COLUNAS}


def exportar_tudo(cat: Catalogo, pastas):
    itens = cat.itens()
    # JSONL (contrato para máquinas): um item por linha, todos os status
    linhas = []
    for it in itens:
        d = {c: it.get(c) for c in COLUNAS + ["processo_tipo", "caminho_origem", "caminho_json", "run_id", "criado_em"]}
        d["campos"] = json.loads(it["campos_json"]) if it.get("campos_json") else {}
        d["contrato_versao"] = 1
        linhas.append(json.dumps(d, ensure_ascii=False))
    escrever_atomico(pastas.jsonl, "\n".join(linhas) + ("\n" if linhas else ""))
    # CSV com ponto e vírgula e UTF-8 com BOM (abre direto no Excel brasileiro)
    P(pastas.csv).parent.mkdir(parents=True, exist_ok=True)
    tmp = pastas.csv + ".tmp"
    with open(P(tmp), "w", encoding="utf-8-sig", newline="") as f:
        w = csv.writer(f, delimiter=";", quoting=csv.QUOTE_MINIMAL)
        w.writerow(COLUNAS)
        for it in itens:
            w.writerow([_celula(it.get(c)) for c in COLUNAS])
    os.replace(P(tmp), P(pastas.csv))
    exportar_xlsx(itens, cat, pastas)


def _celula(v):
    if v is None:
        return ""
    if isinstance(v, float):
        return f"{v:.3f}".replace(".", ",")
    return str(v)


def exportar_xlsx(itens: list[dict], cat: Catalogo, pastas):
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill
    from openpyxl.utils import get_column_letter

    wb = Workbook()
    ws = wb.active
    ws.title = "Resumo"
    por_status: dict[str, int] = {}
    por_cat: dict[str, dict] = {}
    por_arq: dict[str, dict] = {}
    for it in itens:
        por_status[it["status"]] = por_status.get(it["status"], 0) + 1
        if it["status"] in ("AUTO", "REVISAR"):
            c = por_cat.setdefault(it["categoria"] or "?", {"AUTO": 0, "REVISAR": 0, "paginas": 0})
            c[it["status"]] += 1
            c["paginas"] += it["n_paginas"] or 0
        a = por_arq.setdefault(it["arquivo_origem"], {})
        a[it["status"]] = a.get(it["status"], 0) + 1
    ws.append(["CATÁLOGO DO PIPELINE DE ESCANEAMENTO - SEINFRA ILHÉUS"])
    ws.append([f"Gerado em {agora()} (versão {VERSAO})"])
    ws.append([])
    ws.append(["Situação", "Quantidade"])
    for k in ("AUTO", "REVISAR", "SOLTO", "ERRO"):
        ws.append([k, por_status.get(k, 0)])
    ws.append([])
    ws.append(["Categoria", "AUTO", "REVISAR", "Páginas"])
    for k, v in sorted(por_cat.items(), key=lambda kv: -(kv[1]["AUTO"] + kv[1]["REVISAR"])):
        ws.append([k, v["AUTO"], v["REVISAR"], v["paginas"]])
    ws.append([])
    ws.append(["Arquivo de origem", "AUTO", "REVISAR", "SOLTO", "ERRO"])
    for k, v in sorted(por_arq.items()):
        ws.append([k, v.get("AUTO", 0), v.get("REVISAR", 0), v.get("SOLTO", 0), v.get("ERRO", 0)])
    ws["A1"].font = Font(bold=True, size=13)
    ws.column_dimensions["A"].width = 48

    def aba(nome, filtro, cols):
        w = wb.create_sheet(nome)
        w.append(cols)
        for c in w[1]:
            c.font = Font(bold=True, color="FFFFFF")
            c.fill = PatternFill("solid", fgColor="1F4E78")
        for it in itens:
            if filtro(it):
                w.append([_excel(it.get(c)) for c in cols])
        w.freeze_panes = "A2"
        w.auto_filter.ref = w.dimensions
        for i, c in enumerate(cols, 1):
            larg = {"assunto": 50, "objeto": 50, "destinatario": 40, "motivo": 60, "caminho_final": 60,
                    "caminho_texto": 40, "nome_final": 42, "arquivo_origem": 34}.get(c, 14)
            w.column_dimensions[get_column_letter(i)].width = larg

    aba("Documentos", lambda i: i["status"] in ("AUTO", "REVISAR"), COLUNAS)
    aba("Soltos", lambda i: i["status"] == "SOLTO", COLUNAS)
    aba("Revisar", lambda i: i["status"] == "REVISAR", COLUNAS)
    aba("Erros", lambda i: i["status"] == "ERRO", COLUNAS)
    tmp = pastas.xlsx + ".tmp"
    wb.save(P(tmp))
    os.replace(P(tmp), P(pastas.xlsx))


def _excel(v):
    if v is None:
        return None
    if isinstance(v, str) and v[:1] in "=+-@":
        return "'" + v          # evita que o Excel interprete texto como fórmula
    return v
