"""Planilhas FICTÍCIAS extras para o teste do editor de planilhas.
Uso: python3 test/sheet-fixtures.py PASTA_DO_COFRE
"""
import os, sys, io, random
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.chart import BarChart, Reference
from openpyxl.drawing.image import Image as XLImage
from openpyxl.formatting.rule import CellIsRule
from openpyxl.worksheet.datavalidation import DataValidation
from openpyxl.workbook.defined_name import DefinedName
from openpyxl.comments import Comment
from openpyxl.worksheet.table import Table, TableStyleInfo
from PIL import Image, ImageDraw

destino = sys.argv[1] if len(sys.argv) > 1 else "/tmp/cofre-planilhas"
pasta = os.path.join(destino, "Planilhas")
os.makedirs(pasta, exist_ok=True)
random.seed(42)

# ---------- 1) planilha grande (20 mil linhas) ----------
wb = Workbook()
ws = wb.active
ws.title = "Lançamentos"
ws.append(["Nº", "Descrição", "Quantidade", "Valor unitário", "Total", "Data"])
for c in ws[1]:
    c.font = Font(bold=True)
    c.fill = PatternFill("solid", fgColor="DDEBF7")
for i in range(2, 20002):
    ws.append([i - 1, f"Item fictício {i - 1}", random.randint(1, 50), round(random.uniform(1, 500), 2), f"=C{i}*D{i}", 46000 + (i % 365)])
    ws.cell(i, 6).number_format = "dd/mm/yyyy"
    ws.cell(i, 4).number_format = '"R$" #,##0.00'
    ws.cell(i, 5).number_format = '"R$" #,##0.00'
ws.column_dimensions["B"].width = 24
ws.column_dimensions["D"].width = 15
ws.column_dimensions["E"].width = 16
ws.column_dimensions["F"].width = 12
ws.freeze_panes = "A2"
wb.save(os.path.join(pasta, "Grande 20 mil linhas.xlsx"))

# ---------- 2) pasta completa: gráfico, imagem, formatação condicional, validação, nome, comentários ----------
wb = Workbook()
ws = wb.active
ws.title = "Dados"
fino = Side(style="thin", color="808080")
ws.append(["Mês", "Receita", "Despesa", "Saldo"])
meses = ["Janeiro", "Fevereiro", "Março", "Abril", "Maio", "Junho"]
for i, m in enumerate(meses, start=2):
    ws.append([m, 1000 + i * 150, 800 + i * 90, f"=B{i}-C{i}"])
ws.append(["Total", "=SUM(B2:B7)", "=SUM(C2:C7)", "=SUM(D2:D7)"])
for row in ws.iter_rows(min_row=1, max_row=8, max_col=4):
    for c in row:
        c.border = Border(left=fino, right=fino, top=fino, bottom=fino)
for r in range(2, 9):
    for col in (2, 3, 4):
        ws.cell(r, col).number_format = '"R$" #,##0.00'
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF")
    c.fill = PatternFill("solid", fgColor="375623")
    c.alignment = Alignment(horizontal="center")
ws["A8"].font = Font(bold=True)
ws.column_dimensions["A"].width = 14
ws.column_dimensions["B"].width = 14
ws.column_dimensions["C"].width = 14
ws.column_dimensions["D"].width = 14
ws.merge_cells("F1:H1")
ws["F1"] = "Observações (mescladas)"
ws["F1"].alignment = Alignment(horizontal="center")
ws["F2"] = "Taxa"
ws["G2"] = 0.05
ws["G2"].number_format = "0.00%"
ws["F3"] = "Saldo com taxa"
ws["G3"] = "=D8*(1+Taxa)"
ws["F4"] = "Situação"
ws["G4"] = "Aprovado"
ws["A11"] = "Texto longo que transborda para as células vizinhas vazias, como no Excel."
ws["A12"] = "Linha 1\nLinha 2 (quebra de texto)"
ws["A12"].alignment = Alignment(wrap_text=True, vertical="top")
ws.row_dimensions[12].height = 32
# formatação condicional
ws.conditional_formatting.add("D2:D7", CellIsRule(operator="lessThan", formula=["300"], fill=PatternFill("solid", fgColor="F4CCCC")))
# validação de dados
dv = DataValidation(type="list", formula1='"Aprovado,Pendente,Recusado"', allow_blank=True)
dv.prompt = "Escolha uma situação"
dv.promptTitle = "Situação"
ws.add_data_validation(dv)
dv.add("G4")
dv2 = DataValidation(type="whole", operator="between", formula1="0", formula2="100")
ws.add_data_validation(dv2)
dv2.add("B10:B20")
# nome definido
wb.defined_names["Taxa"] = DefinedName("Taxa", attr_text="Dados!$G$2")
# comentários
ws["B2"].comment = Comment("Receita fictícia de janeiro", "Equipe de teste")
ws["D8"].comment = Comment("Soma dos saldos", "Equipe de teste")
# hiperlink
ws["F6"] = "Site de exemplo"
ws["F6"].hyperlink = "https://example.com/"
ws["F6"].font = Font(color="0563C1", underline="single")
# gráfico
graf = BarChart()
graf.title = "Receita x Despesa"
graf.add_data(Reference(ws, min_col=2, max_col=3, min_row=1, max_row=7), titles_from_data=True)
graf.set_categories(Reference(ws, min_col=1, min_row=2, max_row=7))
graf.width = 12
graf.height = 7
ws.add_chart(graf, "J2")
# imagem
img = Image.new("RGB", (120, 60), (31, 78, 121))
d = ImageDraw.Draw(img)
d.rectangle([10, 10, 110, 50], outline=(255, 255, 255), width=3)
buf = io.BytesIO(); img.save(buf, "PNG"); buf.seek(0)
xi = XLImage(buf)
ws.add_image(xi, "J18")
# impressão
ws.print_options.gridLines = True
ws.page_setup.orientation = "landscape"
ws.print_area = "A1:D8"
ws.freeze_panes = "B2"
# segunda aba com tabela e nome com espaço
ws2 = wb.create_sheet("Resumo Geral")
ws2["A1"] = "Saldo total"
ws2["B1"] = "=Dados!D8"
ws2["B1"].number_format = '"R$" #,##0.00'
ws2["A2"] = "Maior receita"
ws2["B2"] = "=MAX(Dados!B2:B7)"
ws2["A4"] = "Código"; ws2["B4"] = "Nome"
for i, n in enumerate(["Alfa", "Beta", "Gama"], start=5):
    ws2.cell(i, 1, i - 4); ws2.cell(i, 2, n)
tab = Table(displayName="TabelaCodigos", ref="A4:B7")
tab.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
ws2.add_table(tab)
ws3 = wb.create_sheet("Oculta")
ws3["A1"] = "Aba oculta de teste"
ws3.sheet_state = "hidden"
wb.save(os.path.join(pasta, "Completa ficticia.xlsx"))

# ---------- 3) CSV com vírgula, BOM, CRLF, aspas e quebras de linha ----------
linhas = [
    ["Nome", "Endereço", "Observação", "Valor"],
    ["Fulano Fictício", "Rua A, 10", 'Disse "olá"', "10.5"],
    ["Beltrana", "Av. B, 200", "Linha 1\nLinha 2", "1234.75"],
    ["Ciclano", "Travessa C", "", "3"],
]
def campo(s):
    return '"' + s.replace('"', '""') + '"' if ("," in s or '"' in s or "\n" in s) else s
txt = "\r\n".join(",".join(campo(c) for c in l) for l in linhas) + "\r\n"
open(os.path.join(pasta, "virgulas-ficticio.csv"), "wb").write(b"\xef\xbb\xbf" + txt.encode("utf-8"))

# ---------- 4) CSV em Windows-1252 (latin1) ----------
txt = "Município;Situação;Valor\r\nSão João Fictício;Ação concluída;1.234,56\r\nVila Imaginária;Em análise;99,9\r\n"
open(os.path.join(pasta, "latin1-ficticio.csv"), "wb").write(txt.encode("cp1252"))
print("planilhas extras criadas em", pasta)
