"""Confere um .docx salvo pelo Tinteiro contra o original (recebido pela entrada padrão na primeira chamada).
Uso: python3 docx_check.py SALVO.docx ORIGINAL.docx   (se o original não existir, é lido da entrada padrão e gravado)
Imprime um JSON na última linha.
"""
import sys, os, json, zipfile, io
from docx import Document

salvo, orig_path = sys.argv[1], sys.argv[2]
if not os.path.exists(orig_path):
    dados = sys.stdin.buffer.read()
    if dados:
        open(orig_path, "wb").write(dados)

res = {"abre": False}
try:
    d = Document(salvo)
    res["abre"] = True
    paragrafos = [p.text for p in d.paragraphs]
    res["paragrafos"] = paragrafos
    res["texto"] = "\n".join(paragrafos)
    res["negrito"] = " ".join(r.text for p in d.paragraphs for r in p.runs if r.bold)
    tab = []
    for t in d.tables:
        for row in t.rows:
            for c in row.cells:
                tab.append(c.text)
    res["tabela"] = " | ".join(tab)
    from docx.enum.text import WD_ALIGN_PARAGRAPH
    res["centralizado"] = [p.text for p in d.paragraphs if p.alignment == WD_ALIGN_PARAGRAPH.CENTER]
    sec = d.sections[0]
    res["cabecalho_ok"] = any("ESTADO FICTÍCIO" in p.text for p in sec.header.paragraphs) and len(sec.header.part.related_parts) > 0
except Exception as e:
    res["erro"] = repr(e)

a, b = zipfile.ZipFile(orig_path), zipfile.ZipFile(salvo)
dif = []
for n in a.namelist():
    if n == "word/document.xml":
        continue
    if n not in b.namelist() or a.read(n) != b.read(n):
        dif.append(n)
res["partes_iguais"] = not dif
res["partes_diferentes"] = dif
res["ordem_igual"] = a.namelist() == b.namelist()
print(json.dumps(res, ensure_ascii=False))
