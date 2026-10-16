"""Gera um .docx fictício com recursos difíceis (para testar o editor de Word).
Uso: python3 docx_complexo.py SAIDA.docx
"""
import sys, io, copy
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.enum.section import WD_ORIENT, WD_SECTION
from docx.oxml.ns import qn
from docx.oxml import OxmlElement
from PIL import Image, ImageDraw

saida = sys.argv[1]
doc = Document()
st = doc.styles["Normal"]; st.font.name = "Arial"; st.font.size = Pt(11)

def hyperlink(par, url, texto):
    part = par.part
    r_id = part.relate_to(url, "http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink", is_external=True)
    h = OxmlElement("w:hyperlink"); h.set(qn("r:id"), r_id)
    r = OxmlElement("w:r"); rpr = OxmlElement("w:rPr")
    c = OxmlElement("w:color"); c.set(qn("w:val"), "0563C1"); rpr.append(c)
    u = OxmlElement("w:u"); u.set(qn("w:val"), "single"); rpr.append(u)
    r.append(rpr); t = OxmlElement("w:t"); t.text = texto; r.append(t); h.append(r)
    par._p.append(h)

def campo(par, instr):
    for tipo in ["begin", None, "separate", "result", "end"]:
        r = OxmlElement("w:r")
        if tipo in ("begin", "separate", "end"):
            f = OxmlElement("w:fldChar"); f.set(qn("w:fldCharType"), tipo); r.append(f)
        elif tipo is None:
            it = OxmlElement("w:instrText"); it.set(qn("xml:space"), "preserve"); it.text = f" {instr} "; r.append(it)
        else:
            t = OxmlElement("w:t"); t.text = "1"; r.append(t)
        par._p.append(r)

def bookmark(par, nome, bid):
    s = OxmlElement("w:bookmarkStart"); s.set(qn("w:id"), str(bid)); s.set(qn("w:name"), nome)
    e = OxmlElement("w:bookmarkEnd"); e.set(qn("w:id"), str(bid))
    par._p.insert(1 if par._p.pPr is not None else 0, s); par._p.append(e)

doc.add_heading("Documento complexo de teste", level=1)
p = doc.add_paragraph("Parágrafo com link: ")
hyperlink(p, "https://example.com/ficticio", "site fictício")
p.add_run(" e texto depois do link.")
p = doc.add_paragraph("Parágrafo com marcador (bookmark) para teste.")
bookmark(p, "marca_teste", 1)
p = doc.add_paragraph("Coluna A\tColuna B\tColuna C com tabulações")
p = doc.add_paragraph("Antes da quebra de página.")
p.add_run().add_break(WD_BREAK.PAGE)
p.add_run("Depois da quebra de página, no mesmo parágrafo.")
# imagem no corpo
im = Image.new("RGB", (200, 80), (230, 240, 255)); d = ImageDraw.Draw(im); d.rectangle([10, 10, 190, 70], outline=(30, 60, 120), width=4)
b = io.BytesIO(); im.save(b, "PNG")
p = doc.add_paragraph("Imagem: "); p.add_run().add_picture(io.BytesIO(b.getvalue()), width=Cm(4))
# tabela com mesclagem
t = doc.add_table(rows=3, cols=3); t.style = "Table Grid"
a = t.cell(0, 0).merge(t.cell(0, 1)); a.text = "Mesclada horizontal"
t.cell(0, 2).text = "C1"
v = t.cell(1, 0).merge(t.cell(2, 0)); v.text = "Mesclada vertical"
t.cell(1, 1).text = "B2"; t.cell(1, 2).text = "C2"; t.cell(2, 1).text = "B3"; t.cell(2, 2).text = "C3"
# tabela aninhada
inner = t.cell(2, 2).add_table(rows=1, cols=2); inner.style = "Table Grid"
inner.cell(0, 0).text = "Interna 1"; inner.cell(0, 1).text = "Interna 2"
doc.add_paragraph("Parágrafo final da primeira seção.")
# seção em paisagem
sec = doc.add_section(WD_SECTION.NEW_PAGE)
sec.orientation = WD_ORIENT.LANDSCAPE; sec.page_width, sec.page_height = sec.page_height, sec.page_width
doc.add_paragraph("Texto na seção em paisagem.")
lista = doc.add_paragraph("Primeiro item", style="List Number")
doc.add_paragraph("Segundo item", style="List Number")
# rodapé com campo de número de página
f = sec.footer.paragraphs[0]; f.text = "Página "
campo(f, "PAGE")
doc.save(saida)
print("ok", saida)
