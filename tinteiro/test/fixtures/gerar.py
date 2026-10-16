"""Gera o cofre de teste com arquivos FICTÍCIOS (nenhum dado real).
Uso: python3 test/fixtures/gerar.py PASTA_DESTINO
"""
import os, sys, io, math, random
from PIL import Image, ImageDraw, ImageFont

destino = sys.argv[1] if len(sys.argv) > 1 else "/tmp/cofre-teste"
os.makedirs(destino, exist_ok=True)
def caminho(*p):
    c = os.path.join(destino, *p)
    os.makedirs(os.path.dirname(c), exist_ok=True)
    return c

# ---------- brasão fictício (escudo desenhado) ----------
def brasao(px=240):
    im = Image.new("RGBA", (px, int(px * 1.15)), (255, 255, 255, 0))
    d = ImageDraw.Draw(im)
    w, h = im.size
    escudo = [(w*0.1, h*0.08), (w*0.9, h*0.08), (w*0.9, h*0.55), (w*0.5, h*0.95), (w*0.1, h*0.55)]
    d.polygon(escudo, fill=(28, 74, 140), outline=(200, 160, 40))
    d.line(escudo + [escudo[0]], fill=(200, 160, 40), width=max(3, px // 40))
    d.rectangle([w*0.42, h*0.15, w*0.58, h*0.75], fill=(240, 240, 240))
    d.rectangle([w*0.22, h*0.33, w*0.78, h*0.45], fill=(240, 240, 240))
    d.ellipse([w*0.40, h*0.0, w*0.60, h*0.12], fill=(200, 160, 40))
    return im
b = brasao()
buf = io.BytesIO(); b.save(buf, "PNG"); png_brasao = buf.getvalue()
open(caminho("Imagens", "brasao-ficticio.png"), "wb").write(png_brasao)

# imagem de paisagem fictícia
im = Image.new("RGB", (900, 560), (210, 230, 250))
d = ImageDraw.Draw(im)
for i in range(0, 560, 4):
    d.line([(0, i), (900, i)], fill=(120 + i // 6, 170 + i // 10, 230 - i // 8))
d.polygon([(0, 560), (250, 260), (520, 560)], fill=(70, 110, 80))
d.polygon([(300, 560), (620, 210), (900, 560)], fill=(50, 90, 70))
d.ellipse([700, 60, 800, 160], fill=(255, 220, 120))
im.save(caminho("Imagens", "paisagem.jpg"), quality=88)

# ---------- DOCX ----------
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

doc = Document()
sec = doc.sections[0]
sec.page_height = Cm(29.7); sec.page_width = Cm(21.0)
sec.top_margin = Cm(3); sec.bottom_margin = Cm(2); sec.left_margin = Cm(3); sec.right_margin = Cm(2)
st = doc.styles["Normal"]
st.font.name = "Times New Roman"; st.font.size = Pt(12)
st.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
pf = st.paragraph_format; pf.space_after = Pt(6); pf.line_spacing = 1.5

# cabeçalho com brasão
hdr = sec.header
hp = hdr.paragraphs[0]
hp.alignment = WD_ALIGN_PARAGRAPH.CENTER
hp.add_run().add_picture(io.BytesIO(png_brasao), width=Cm(2.0))
h2 = hdr.add_paragraph(); h2.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h2.add_run("ESTADO FICTÍCIO DE EXEMPLO"); r.bold = True; r.font.size = Pt(11); r.font.name = "Arial"
h3 = hdr.add_paragraph(); h3.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = h3.add_run("Prefeitura Municipal de Cidade Imaginária — Secretaria de Testes"); r.font.size = Pt(10); r.font.name = "Arial"
ftr = sec.footer.paragraphs[0]; ftr.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = ftr.add_run("Rua das Flores, 123 — Centro — CEP 00000-000 — Documento fictício para testes"); r.font.size = Pt(8); r.font.name = "Arial"

t = doc.add_paragraph(); t.alignment = WD_ALIGN_PARAGRAPH.CENTER
r = t.add_run("TERMO DE REFERÊNCIA Nº 001/2026"); r.bold = True; r.font.size = Pt(14)
doc.add_heading("1. Do objeto", level=1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.paragraph_format.first_line_indent = Cm(1.25)
p.add_run("O presente termo tem por objeto a ")
r = p.add_run("contratação fictícia"); r.bold = True
p.add_run(" de serviços de manutenção de praças imaginárias, conforme as especificações a seguir. ")
r = p.add_run("Este texto existe apenas para testes."); r.italic = True
doc.add_heading("2. Das especificações", level=1)
p = doc.add_paragraph("Item com marcador A", style="List Bullet")
p = doc.add_paragraph("Item com marcador B", style="List Bullet")
p = doc.add_paragraph("Primeiro passo numerado", style="List Number")
p = doc.add_paragraph("Segundo passo numerado", style="List Number")

tab = doc.add_table(rows=4, cols=4); tab.style = "Table Grid"; tab.alignment = WD_TABLE_ALIGNMENT.CENTER
cab = ["Item", "Descrição", "Qtd.", "Valor (R$)"]
linhas = [["1", "Poda de árvores fictícias", "10", "1.500,00"], ["2", "Pintura de bancos imaginários", "25", "3.250,50"], ["3", "Limpeza de lago de exemplo", "1", "980,00"]]
for i, c in enumerate(cab):
    cel = tab.rows[0].cells[i]; cel.text = ""
    rr = cel.paragraphs[0].add_run(c); rr.bold = True
    shd = OxmlElement("w:shd"); shd.set(qn("w:val"), "clear"); shd.set(qn("w:color"), "auto"); shd.set(qn("w:fill"), "D9E2F3")
    cel._tc.get_or_add_tcPr().append(shd)
for li, lin in enumerate(linhas):
    for ci, v in enumerate(lin):
        tab.rows[li + 1].cells[ci].text = v
        if ci >= 2: tab.rows[li + 1].cells[ci].paragraphs[0].alignment = WD_ALIGN_PARAGRAPH.RIGHT
doc.add_paragraph()
doc.add_heading("3. Do prazo", level=1)
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
p.add_run("O prazo de execução é de 30 (trinta) dias corridos, contados da emissão da ordem de serviço fictícia. ")
r = p.add_run("Texto sublinhado de exemplo"); r.underline = True
p.add_run(" e texto ")
r = p.add_run("em vermelho"); r.font.color.rgb = RGBColor(0xC0, 0x00, 0x00)
p.add_run(".")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("Cidade Imaginária, 16 de outubro de 2026.")
p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
p.add_run("\n\nNome do Signatário\n").bold = True
p.add_run("Cargo do Signatário")
doc.save(caminho("Documentos", "Termo de Referencia ficticio.docx"))

# docx simples, segundo arquivo
d2 = Document()
d2.add_heading("Ofício de teste", 0)
d2.add_paragraph("Parágrafo um do ofício fictício.")
d2.add_paragraph("Parágrafo dois com mais palavras para a busca: abacaxi azul.")
d2.save(caminho("Documentos", "Oficio teste.docx"))

# arquivo com @ (o app abre normalmente)
d3 = Document(); d3.add_paragraph("Arquivo fictício cujo nome começa com arroba.")
d3.save(caminho("Documentos", "@Arroba ficticio.docx"))

# ---------- XLSX ----------
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
wb = Workbook()
ws = wb.active; ws.title = "Orçamento"
fino = Side(style="thin", color="999999")
ws.append(["Item", "Descrição", "Qtd", "Unitário", "Total"])
for c in ws[1]:
    c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="1F4E79"); c.alignment = Alignment(horizontal="center")
dados = [(1, "Cimento fictício", 10, 35.5), (2, "Areia imaginária", 4, 120), (3, "Tinta de teste", 7, 89.9), (4, "Pincel exemplo", 12, 15.25)]
for i, (n, desc, q, u) in enumerate(dados, start=2):
    ws.append([n, desc, q, u, f"=C{i}*D{i}"])
ws.append(["", "TOTAL", "=SUM(C2:C5)", "", "=SUM(E2:E5)"])
for row in ws.iter_rows(min_row=1, max_row=6, max_col=5):
    for c in row: c.border = Border(left=fino, right=fino, top=fino, bottom=fino)
for r in range(2, 7):
    ws.cell(r, 4).number_format = '"R$" #,##0.00'; ws.cell(r, 5).number_format = '"R$" #,##0.00'
ws["B6"].font = Font(bold=True); ws["E6"].font = Font(bold=True)
ws.column_dimensions["B"].width = 28; ws.column_dimensions["D"].width = 14; ws.column_dimensions["E"].width = 16
ws.freeze_panes = "A2"
ws2 = wb.create_sheet("Resumo")
ws2["A1"] = "Total geral"; ws2["B1"] = "=Orçamento!E6"; ws2["B1"].number_format = '"R$" #,##0.00'
ws2["A2"] = "Média unitária"; ws2["B2"] = "=AVERAGE(Orçamento!D2:D5)"
ws2["A3"] = "Data"; ws2["B3"] = "=DATE(2026,10,16)"; ws2["B3"].number_format = "dd/mm/yyyy"
ws2.merge_cells("A5:C5"); ws2["A5"] = "Células mescladas de exemplo"; ws2["A5"].alignment = Alignment(horizontal="center")
wb.save(caminho("Planilhas", "Orcamento ficticio.xlsx"))

open(caminho("Planilhas", "contatos-ficticios.csv"), "w", encoding="utf-8").write(
    "Nome;Cidade;Valor\nFulano de Tal;Cidade Imaginária;100\nBeltrana Exemplo;Vila Teste;250,5\nCiclano Fictício;Porto Modelo;75\n")

# ---------- PDF ----------
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas as rlcanvas
from reportlab.lib.utils import ImageReader
def pdf(nome, paginas, titulo):
    c = rlcanvas.Canvas(caminho("PDFs", nome), pagesize=A4)
    w, h = A4
    for i in range(1, paginas + 1):
        c.drawImage(ImageReader(io.BytesIO(png_brasao)), w/2 - 25, h - 95, width=50, height=57, mask="auto")
        c.setFont("Helvetica-Bold", 16); c.drawCentredString(w/2, h - 120, f"{titulo}")
        c.setFont("Helvetica", 12); c.drawCentredString(w/2, h - 140, f"Página {i} de {paginas} — documento fictício")
        c.setFont("Helvetica", 11)
        y = h - 190
        for linha in range(18):
            c.drawString(72, y, f"Linha {linha+1} da página {i}: texto de exemplo para testar seleção e busca (laranja).")
            y -= 20
        c.setStrokeColorRGB(0.6, 0.6, 0.6); c.rect(72, 120, w - 144, 120)
        c.drawString(80, 220, "Quadro para assinatura (fictícia)")
        c.showPage()
    c.save()
pdf("Relatorio ficticio.pdf", 4, "RELATÓRIO FICTÍCIO DE TESTE")
pdf("Anexo ficticio.pdf", 2, "ANEXO FICTÍCIO")

# ---------- Markdown / TXT / HTML ----------
open(caminho("Bem-vindo.md"), "w", encoding="utf-8").write("""# Bem-vindo ao Tinteiro

Este é um **cofre de teste** com arquivos *fictícios*. Veja também [[Notas/Lista de tarefas]].

## Recursos
- Pré-visualização ao vivo, como num caderno digital
- Links internos com `[[colchetes duplos]]`
- Listas de tarefas, tabelas e citações

> [!note] Observação
> Os documentos desta pasta são inventados.

### Tabela

| Coluna A | Coluna B |
|---|---|
| um | dois |
| três | quatro |

```js
console.log("olá, mundo");
```

Fim do exemplo com a palavra abacaxi.
""")
open(caminho("Notas", "Lista de tarefas.md"), "w", encoding="utf-8").write("""# Lista de tarefas

- [x] Instalar o app
- [ ] Abrir um documento Word
- [ ] Editar uma planilha
- [ ] Assinar um PDF

Voltar para [[Bem-vindo]].
""")
open(caminho("Notas", "Rascunho.txt"), "w", encoding="utf-8").write("Rascunho em texto puro.\nSegunda linha do rascunho.\n")
open(caminho("Relatorios", "Relatorio exemplo.html"), "w", encoding="utf-8").write("""<!doctype html><html lang="pt-BR"><head><meta charset="utf-8"><title>Relatório de exemplo</title>
<style>body{font-family:Georgia,serif;max-width:720px;margin:40px auto;padding:0 16px;color:#222}h1{color:#1f4e79}table{border-collapse:collapse}td,th{border:1px solid #bbb;padding:6px 10px}</style></head>
<body><h1>Relatório fictício</h1><p>Este relatório HTML é renderizado dentro do app.</p><img src="grafico.png" alt="gráfico" width="300">
<table><tr><th>Mês</th><th>Valor</th></tr><tr><td>Janeiro</td><td>10</td></tr><tr><td>Fevereiro</td><td>20</td></tr></table>
<p><button onclick="this.textContent='Clicado!'">Botão interativo</button></p></body></html>""")
g = Image.new("RGB", (300, 160), "white"); gd = ImageDraw.Draw(g)
for i, v in enumerate([40, 90, 60, 130, 100]):
    gd.rectangle([20 + i * 55, 150 - v, 60 + i * 55, 150], fill=(31, 78, 121))
g.save(caminho("Relatorios", "grafico.png"))
open(caminho("Outros", "apresentacao-ficticia.pptx"), "wb").write("arquivo ficticio (formato sem editor no app)".encode())
os.makedirs(caminho("Pasta vazia", ".keep").rsplit(os.sep, 1)[0], exist_ok=True)
print("cofre de teste criado em", destino)
